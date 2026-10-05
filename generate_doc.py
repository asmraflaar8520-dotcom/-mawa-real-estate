from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Add Title
title = doc.add_heading('دليل البنية التحتية ومعمارية مشروع مأوى (Mawa Architecture Guide)', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 1. مقدمة عن المعمارية
doc.add_heading('1. المعمارية العامة للمشروع (Project Architecture)', level=1)
doc.add_paragraph("مشروع 'مأوى' هو منصة عقارية تعمل بمعمارية فك الارتباط (Decoupled Architecture) بين الواجهة الأمامية (Frontend) والخلفية (Backend)، ولكن يتم استضافتهما معاً على منصة Vercel كخدمة واحدة (Monorepo-style deployment).")
doc.add_paragraph("التقنيات المستخدمة:")
doc.add_paragraph("• الواجهة الأمامية (Frontend): HTML5, CSS3 (Vanilla), JavaScript (Vanilla).", style='List Bullet')
doc.add_paragraph("• الواجهة الخلفية (Backend): Python باستخدام إطار عمل FastAPI.", style='List Bullet')
doc.add_paragraph("• قاعدة البيانات (Database): SQLite محلياً (يتم تفريغه دورياً على Vercel Serverless لأنه Ephemeral) للبيانات السريعة.", style='List Bullet')
doc.add_paragraph("• الاستضافة (Hosting): Vercel (Static Hosting للواجهة الأمامية + Serverless Functions للواجهة الخلفية).", style='List Bullet')

# 2. الهيكل التنظيمي
doc.add_heading('2. هيكل المجلدات والملفات (Folder Structure)', level=1)
p = doc.add_paragraph()
p.add_run("• /backend: ").bold = True
p.add_run("يحتوي على كافة ملفات الواجهة الخلفية (FastAPI).\n")
p.add_run("  - /app/main.py: نقطة الدخول (Entry point) ويحتوي على إعدادات الأمان (CORS, Security Headers).\n")
p.add_run("  - /app/api/: يحتوي على المسارات (Endpoints) مثل auth.py (التسجيل والدخول) و properties.py (إدارة العقارات).\n")
p.add_run("  - /app/models/: جداول قاعدة البيانات (SQLAlchemy).\n")
p.add_run("  - /app/schemas/: هياكل التحقق من البيانات (Pydantic).\n")

p = doc.add_paragraph()
p.add_run("• /public (أو /frontend): ").bold = True
p.add_run("يحتوي على ملفات الواجهة الأمامية.\n")
p.add_run("  - /js/app.js: الملف الرئيسي للبرمجة النصية في المتصفح.\n")
p.add_run("  - /css/design-system.css: نظام التصميم والتنسيقات.\n")
p.add_run("  - ملفات HTML: index.html, property.html, compare.html وغيرها.")

p = doc.add_paragraph()
p.add_run("• vercel.json: ").bold = True
p.add_run("ملف إعدادات النشر على Vercel ويوجه مسار /api إلى FastAPI ويضع ترويسات الأمان (Security Headers).")

# 3. رحلة المستخدم والدوال المرتبطة
doc.add_heading('3. رحلة المستخدم (User Journey) ووظائف الدوال (Functions)', level=1)

# تصفح العقارات
doc.add_heading('أ. تصفح والبحث عن العقارات', level=2)
doc.add_paragraph("عند دخول المستخدم للصفحة الرئيسية (index.html)، يتم تحميل الواجهة واستدعاء دالة fetchProperties() من ملف app.js. هذه الدالة ترسل طلب GET إلى المسار /api/properties.")
doc.add_paragraph("في Backend (properties.py)، الدالة get_properties() تتصل بقاعدة البيانات لجلب العقارات وإرجاعها كـ JSON. بعدها في Frontend، دالة renderProperties(properties) تقوم بتكوين الـ HTML وعرضها للمستخدم، مع استخدام App.escapeHtml() لمنع ثغرات XSS.")

# تسجيل الدخول
doc.add_heading('ب. تسجيل الدخول والتسجيل (Authentication)', level=2)
doc.add_paragraph("عند تعبئة نموذج الدخول (login.html) يتم استدعاء دالة Login التي ترسل طلب POST إلى /api/auth/login. في Backend، الدالة تتحقق من كلمة المرور المشفرة (Bcrypt) وتعيد رمز JWT (JSON Web Token). يتم حفظ هذا الرمز في المتصفح (localStorage/sessionStorage) لتعريف المستخدم في الطلبات القادمة.")

# عرض التفاصيل
doc.add_heading('ج. عرض تفاصيل عقار', level=2)
doc.add_paragraph("عند الضغط على عقار، ينتقل المستخدم إلى property.html?id=X. يتم قراءة الـ id من الرابط واستدعاء /api/properties/{id} من خلال الـ backend لإرجاع تفاصيل العقار كاملة وعرضها.")

# المقارنة
doc.add_heading('د. مقارنة العقارات', level=2)
doc.add_paragraph("يوجد نظام مقارنة محلي (Compare) يعتمد على تخزين معرّفات العقارات (IDs) في localStorage. عند الضغط على (أضف للمقارنة)، تقوم الدالة بحفظ المعرف. وفي صفحة compare.html، يتم جلب بيانات هذه العقارات لعرضها في جدول.")

# 4. التخزين وقواعد البيانات
doc.add_heading('4. أين يتم تخزين البيانات؟', level=1)
doc.add_paragraph("• بيانات المستخدمين والعقارات: تُخزن في قاعدة بيانات SQLite محلياً (ملف mawa.db في مجلد backend).", style='List Bullet')
doc.add_paragraph("• الجلسات (Sessions): لا يوجد تخزين جلسات في السيرفر (Stateless). نعتمد على JWT لتوثيق المستخدمين.", style='List Bullet')
doc.add_paragraph("• بيانات المقارنة (Compare): تُخزن في متصفح المستخدم (localStorage) لسرعة الوصول.", style='List Bullet')
doc.add_paragraph("• الصور والوسائط: إما روابط خارجية أو مخزنة محلياً في مجلد الأصول.", style='List Bullet')

# 5. الأمان
doc.add_heading('5. إجراءات الأمان المطبقة', level=1)
doc.add_paragraph("تم إضافة نظام حماية متكامل (Security Headers) عبر Middleware في FastAPI وعبر vercel.json:")
doc.add_paragraph("• CSP (Content Security Policy): يمنع تنفيذ أي سكربتات خارجية خبيثة (XSS).", style='List Bullet')
doc.add_paragraph("• X-Frame-Options: يمنع تضمين الموقع في إطارات خارجية لحماية من Clickjacking.", style='List Bullet')
doc.add_paragraph("• HSTS: يجبر المتصفحات على استخدام اتصال HTTPS مشفر فقط.", style='List Bullet')
doc.add_paragraph("• دالة App.escapeHtml() في app.js: تعقم أي نصوص يتم إدخالها في HTML لمنع حقن السكربتات.", style='List Bullet')

# Save Document
doc.save('w:/مأوي/mawa/Mawa_Architecture_Guide.docx')
